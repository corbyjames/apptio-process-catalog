# -*- coding: utf-8 -*-
"""Excel workbook generator — called by generate.py. Reads the catalog.json shape."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864"); HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BASE = Font(name=FONT, size=10); BOLD = Font(name=FONT, size=10, bold=True)
THIN = Border(*[Side(style="thin", color="BFBFBF")]*4); WRAP = Alignment(wrap_text=True, vertical="top")
L0_FILLS = {"01":"DEEBF7","02":"DEEBF7","03":"DEEBF7","04":"E2EFDA","05":"FFF2CC","06":"FCE4D6","07":"E4DFEC","08":"FCE4D6","10":"EDEDED"}
FLOW_FILLS = {"UC3":"DEEBF7","UC2":"E2EFDA","UC1":"FFF2CC","UC4":"FCE4D6","CLD":"E4DFEC","INV":"D9D9D9","ZBB":"FBE5D6"}

def build(ROOT, CAT, L0S, L1S, L2S, FLOWS, FW, CFG, CAL, DIAGRAMS):
    wb = Workbook(); V = CAT["version"]
    def header(ws, cols, widths):
        for i, (c, w) in enumerate(zip(cols, widths), 1):
            cell = ws.cell(1, i, c); cell.font = HDR_FONT; cell.fill = HDR_FILL; cell.alignment = WRAP; cell.border = THIN
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"
    def put(ws, r, vals, fill=None, bold_col=None):
        for j, val in enumerate(vals, 1):
            c = ws.cell(r, j, val); c.font = BASE; c.alignment = WRAP; c.border = THIN
        if fill: ws.cell(r, 1).fill = PatternFill("solid", fgColor=fill)
        if bold_col: ws.cell(r, bold_col).font = BOLD

    ws = wb.active; ws.title = "Read Me"; ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 4; ws.column_dimensions["B"].width = 120
    rows = [
    (f"IBM Apptio Process Catalog v{V} - Targetprocess, Costing, Planning, Cloudability", True),
    (f"Generated {CAT['built']} from data/catalog.json in github.com/corbyjames/apptio-process-catalog. Structure: {len(L0S)} L0 areas · {len(L1S)} L1 groups · {len(L2S)} L2 processes · {len(FLOWS)} cross-tool flows · {len(DIAGRAMS)} BPMN diagrams.", False),
    ("", False),
    ("Purpose: a single catalog of the business processes the four IBM Apptio products support and enable, framed by TBM, SPM (hybrid and agile) and FinOps, structured to generate L0-L2 BPMN diagrams and to drive tool configuration.", False),
    ("", False),
    ("How the levels work", True),
    ("L0 = process area. L1 = process group. L2 = process/activity - the unit that becomes a BPMN task. Positional IDs (NN, NN.N, NN.N.N) are the human-facing reference; permanent element IDs (P-0001..., G-011...) survive re-parenting.", False),
    ("Delivery Model: 'Any' = methodology-agnostic; 'Agile' = SAFe/agile-specific; 'Traditional' = waterfall/stage-gate; 'Hybrid' = explicitly about running both side by side. Budgeting method (02.4, 05.1, 05.8, 07.6): Incremental, Driver-based, ZBB, Rolling, Lean/participatory, Any.", False),
    ("BPMN lane / task type: derived per L2 from the Who text and the area's default lanes; these drive the generated group (NN.N.bpmn) and area (L0-NN.bpmn) diagrams. Flow diagrams (flow-XXX.bpmn) come from the E2E Flow Steps sheet.", False),
    ("", False),
    ("Sheets", True),
    ("L0 Map - the nine process areas with tool coverage, default BPMN lanes and counts.", False),
    ("Process Catalog L0-L2 - the full catalog, one row per L2 with element ID, lane, task type, flows, personas and cadence.", False),
    ("E2E Flow Steps - BPMN-ready step tables for the seven cross-tool flows (UC3, UC2, UC1, UC4, CLD, INV, ZBB).", False),
    ("BPMN Diagrams - index of every generated diagram (file names in diagrams/bpmn/generated and assets/diagrams).", False),
    ("Tool Config Reference - configuration-object checklists per product.", False),
    ("Framework Reference - TBM, FinOps, SPM, SAFe structures and maturity scales.", False),
    ("Operating Calendar - fiscal-month rhythm of the catalog's processes.", False),
    ("Changelog - versions of the catalog.", False),
    ]
    for i, (t, b) in enumerate(rows, 2):
        c = ws.cell(i, 2, t); c.font = Font(name=FONT, size=14 if i == 2 else 10, bold=b); c.alignment = WRAP

    ws = wb.create_sheet("L0 Map")
    header(ws, ["L0 ID","Process Area","Description","Primary Product","Band","Default BPMN Lanes","# L1","# L2"], [7,30,60,16,10,45,7,7])
    for i, l0 in enumerate(L0S, 2):
        put(ws, i, [l0["id"], l0["name"], l0["description"], l0["primary"], l0["band"], "; ".join(l0["lanes"]), len(l0["l1s"]), sum(len(g["l2s"]) for g in l0["l1s"])], L0_FILLS.get(l0["id"]), 2)

    ws = wb.create_sheet("Process Catalog L0-L2")
    cols = ["L0 ID","L0 Area","L1 ID","L1 Group","L1 EID","L2 ID","L2 Process","Element ID","Legacy ID","What you get","Description","Delivery Model","Budgeting Method","Product (tool)","Primary Product","Personas (source)","Persona roles","BPMN Lane","BPMN Task Type","Cadence","Cadence bucket","Inputs","Outputs","Configuration Objects","Framework Mapping","Framework tags","Flows","Evidence"]
    header(ws, cols, [6,22,7,26,7,8,32,8,8,30,50,10,16,18,12,24,26,18,12,16,12,26,26,50,34,12,10,30])
    r = 2
    l0name = {l0["id"]: l0["name"] for l0 in L0S}; l1name = {g["id"]: g for g in L1S}
    for l2 in L2S:
        g = l1name[l2["l1"]]
        put(ws, r, [l2["l0"], l0name[l2["l0"]], g["id"], g["name"], g["eid"], l2["id"], l2["name"], l2["eid"], l2.get("legacy_id") or "", l2["outcome"], l2["description"], l2["delivery_model"], l2.get("budgeting_method") or "", l2["tool"], l2["primary_product"], l2["personas"], ", ".join(l2["persona_list"]), l2.get("lane",""), l2.get("bpmn_type",""), l2["cadence"], l2["cadence_bucket"], l2["inputs"], l2["outputs"], l2["config"], l2["framework"], ", ".join(l2["framework_tags"]), ", ".join(l2["flows"]), l2["evidence"]], L0_FILLS.get(l2["l0"]), 7)
        r += 1

    ws = wb.create_sheet("E2E Flow Steps")
    header(ws, ["Flow","Flow Name","Title","Step","BPMN Element Type","Lane","Task / Event","Catalog L2s","Notes"], [7,18,30,6,16,26,55,18,45])
    r = 2
    for f in FLOWS:
        for s in f["steps"]:
            put(ws, r, [f["id"], f["name"], f["title"], s["n"], s["type"], s["lane"], s["task"], ", ".join(s["l2"]), s.get("note","") or ""], FLOW_FILLS.get(f["id"])); r += 1

    ws = wb.create_sheet("BPMN Diagrams")
    header(ws, ["File base name","Kind","Name","Lanes","Steps","BPMN file","SVG file"], [14,8,44,50,7,40,34])
    for i, (base, f) in enumerate(DIAGRAMS, 2):
        kind = "flow" if not f.get("kind") else ("area" if f["kind"] == "l0" else "group")
        put(ws, i, [base, kind, f["name"] if f.get("kind") else f["id"]+" "+f["name"], "; ".join(f["lanes"]), len(f["steps"]), f"diagrams/bpmn/generated/{base}.bpmn", f"assets/diagrams/{base}.svg"], bold_col=1)

    ws = wb.create_sheet("Tool Config Reference")
    header(ws, ["Tool","Config Domain","Configuration Objects","Notes"], [16,28,70,45])
    for i, c in enumerate(CFG, 2): put(ws, i, [c["tool"], c["domain"], c["objects"], c.get("notes") or ""], bold_col=1)

    ws = wb.create_sheet("Framework Reference")
    header(ws, ["Framework","Structure / Layers","Disciplines / Capabilities","Maturity & Personas"], [12,55,70,55])
    for i, (k, v) in enumerate(FW.items(), 2): put(ws, i, [k, v.get("layers",""), v.get("disciplines",""), v.get("maturity","")], bold_col=1)

    ws = wb.create_sheet("Operating Calendar")
    header(ws, ["Stream","Process","Cadence"] + ["FM%d" % i for i in range(1,13)] + ["Description","Catalog refs"], [22,34,15]+[5]*12+[60,22])
    r = 2
    for st in CAL["streams"]:
        for en in st["entries"]:
            put(ws, r, [st["name"], en["name"], en["cadence"]] + [("X" if m in en["fm"] else "") for m in range(1,13)] + [en["desc"], ", ".join(en.get("refs", []))]); r += 1

    ws = wb.create_sheet("Changelog")
    header(ws, ["Version","Date","Change"], [9,12,120])
    r = 2
    for rel in CAT["changelog"]:
        for it in rel["items"]: put(ws, r, [rel["version"], rel["date"], it]); r += 1

    wb.save(os.path.join(ROOT, "dist", "Apptio_Process_Catalog_L0-L2.xlsx"))
